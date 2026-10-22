"""独立阅读笔记：Markdown 存储、批注来源与双链解析，显式保存才落盘。"""

from __future__ import annotations

import hashlib
import json
import logging
import re
from collections.abc import Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.parse import quote

from pdfreader.engines.types import Annotation

TEMPLATES = {
    "期刊论文": ("研究问题", "方法与数据", "主要结果", "局限", "我的思考"),
    "学位论文": ("研究背景", "章节提纲", "核心贡献", "方法细节", "待核对问题"),
    "综述": ("范围与检索", "分类框架", "关键证据", "争议与空白", "后续阅读"),
    "实验": ("目的与假设", "材料与步骤", "观察与结果", "异常与解释", "下一步"),
}
_SOURCE = re.compile(r"<!-- lumen-source:(\{[^\n]*\}) -->")
_WIKILINK = re.compile(r"\[\[([^\]\n]+)\]\]")


@dataclass(frozen=True)
class SourceAnchor:
    fingerprint: str
    page: int
    quote: str
    document_key: str = ""


def annotation_fingerprint(annotation: Annotation, document_key: str = "") -> str:
    """不使用易重排的页内索引；同源歧义时由解析器报告，避免跳到错误批注。"""
    payload = [
        document_key,
        annotation.stable_id,
        annotation.annot_type,
        [
            round(v, 2)
            for v in (
                annotation.rect.x0,
                annotation.rect.y0,
                annotation.rect.x1,
                annotation.rect.y1,
            )
        ],
        annotation.marked_text,
        annotation.content,
    ]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False).encode()).hexdigest()


def template(title: str, kind: str) -> str:
    return "# " + title + "\n\n" + "\n\n".join("## " + h + "\n" for h in TEMPLATES[kind]) + "\n"


def annotation_note(annotation: Annotation, document_key: str = "") -> str:
    anchor = SourceAnchor(
        annotation_fingerprint(annotation, document_key),
        annotation.page_num,
        annotation.marked_text,
        document_key,
    )
    metadata = json.dumps(asdict(anchor), ensure_ascii=False).replace("-->", "--\\u003e")
    quoted = "\n".join("> " + line for line in annotation.marked_text.splitlines())
    return (
        f"\n<!-- lumen-source:{metadata} -->\n"
        f"[原文 · 第 {annotation.page_num + 1} 页](source:{anchor.fingerprint})\n\n"
        f"{quoted}\n\n{annotation.content}\n"
    )


def source_anchors(markdown: str) -> list[SourceAnchor]:
    result: list[SourceAnchor] = []
    for match in _SOURCE.finditer(markdown):
        try:
            data = json.loads(match.group(1))
            fingerprint, page, text = data["fingerprint"], data["page"], data["quote"]
            if isinstance(fingerprint, str) and isinstance(page, int) and isinstance(text, str):
                document_key = data.get("document_key", "")
                if isinstance(document_key, str):
                    result.append(SourceAnchor(fingerprint, page, text, document_key))
        except (ValueError, KeyError, TypeError):
            # 用户可自由编辑 Markdown；损坏的来源保留原文，由界面显示不可解析。
            continue
    return result


def display_markdown(markdown: str) -> str:
    """来源元数据随文件保存，但不占用用户编辑区。"""
    return re.sub(r"<!-- lumen-source:[^\n]* -->\n?", "", markdown)


def annotated_markdown(markdown: str, anchors: Sequence[SourceAnchor]) -> str:
    active = [anchor for anchor in anchors if f"source:{anchor.fingerprint}" in markdown]
    metadata = [
        "<!-- lumen-source:"
        + json.dumps(asdict(anchor), ensure_ascii=False).replace("-->", "--\\u003e")
        + " -->"
        for anchor in active
    ]
    return ("\n".join(metadata) + "\n" if metadata else "") + markdown


def resolve_sources(
    anchors: Sequence[SourceAnchor], annotations: Sequence[Annotation], document_key: str = ""
) -> dict[str, Annotation | None]:
    grouped: dict[str, list[Annotation]] = {}
    for annotation in annotations:
        grouped.setdefault(annotation_fingerprint(annotation, document_key), []).append(annotation)
    return {
        anchor.fingerprint: matches[0]
        if anchor.document_key == document_key
        and len(matches := grouped.get(anchor.fingerprint, [])) == 1
        else None
        for anchor in anchors
    }


def preview_markdown(markdown: str) -> str:
    return _WIKILINK.sub(lambda m: f"[{m[1]}](note:{quote(m[1], safe='')})", markdown)


def linked_note_path(directory: Path, title: str) -> Path:
    """双链只允许同一笔记目录中的普通 Markdown 文件名。"""
    if not title or any(c in title for c in "/\\:\x00") or title in (".", ".."):
        raise ValueError("双链名称必须是同目录的笔记名")
    target = (directory / (title + ".md")).resolve()
    if target.parent != directory.resolve():
        raise ValueError("双链不能跳出笔记目录")
    return target


@dataclass(frozen=True)
class Backlinks:
    paths: tuple[Path, ...]
    failed: tuple[Path, ...]
    truncated: bool = False


def find_backlinks(directory: Path, title: str) -> Backlinks:
    paths = sorted(
        path for path in directory.iterdir() if path.suffix.lower() == ".md" and path.is_file()
    )
    found: list[Path] = []
    failed: list[Path] = []
    for path in paths[:1000]:
        if path.stem.casefold() == title.casefold() or path.resolve().parent != directory.resolve():
            continue
        try:
            if path.stat().st_size > 2_000_000:
                raise ValueError("笔记超过反向链接扫描上限")
            targets = [match[1].casefold() for match in _WIKILINK.finditer(path.read_text("utf-8"))]
            if title.casefold() in targets:
                found.append(path)
        except (OSError, UnicodeError, ValueError):
            logging.getLogger("pdfreader").warning("反向链接扫描失败：%s", path, exc_info=True)
            failed.append(path)
    return Backlinks(tuple(found), tuple(failed), len(paths) > 1000)


def save_markdown(path: Path, markdown: str) -> None:
    """先写完整临时文件再替换目标；写入失败不覆盖原笔记。"""
    from uuid import uuid4

    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    temporary.write_text(markdown, encoding="utf-8")
    temporary.replace(path)


def export_docx(path: Path, markdown: str) -> None:
    """笔记独立导出，不将 PDF→DOCX 转换器误当成 Markdown 写入器。"""
    from uuid import uuid4

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.element.get_or_add_rPr().append(fonts)
    for line in _SOURCE.sub("", markdown).splitlines():
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            document.add_heading(heading[2], level=len(heading[1]))
        else:
            document.add_paragraph(line)
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    document.save(str(temporary))
    temporary.replace(path)
